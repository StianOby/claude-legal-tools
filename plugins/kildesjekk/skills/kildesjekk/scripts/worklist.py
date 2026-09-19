#!/usr/bin/env python3
"""worklist.py - check that a kildesjekk worklist is well formed.

    python worklist.py validate Kildesjekk_<documentname>.xlsx [--final] [--json]

This reads the worklist and reports every row that breaks one of the
invariants in SKILL.md ("Row invariants"), plus the status counts that §16
asks for. It never writes to the file.

What it can and cannot tell you: it checks the *shape* of the worklist -
that the permitted values are used, that a "source unavailable" row carries
no severity, that no two rows share a key, that the colours match the state,
that the Metadata counts add up. It cannot tell you whether a source was
really consulted or a quotation really compared. A clean report means the
worklist is well formed, not that the source check is correct.

Exit status: 0 when no violations were found, 1 when there were, 2 when the
file could not be read as a worklist at all.
"""

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - environment-dependent
    sys.exit("worklist.py needs openpyxl: pip install openpyxl")

# The report quotes the worklist back - references, quotations, Norwegian
# and other non-ASCII text. A console that cannot encode a character must
# not take the whole report down with it.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

CHECKED_VALUES = ("yes", "no", "source unavailable")
SEVERITIES = ("high", "medium", "low")

# Each column is found by its header: the exact one first, then by keyword,
# so that a reworded header still matches. "discrepancy" and "type of
# discrepancy" both contain the word, and "reference" appears in two
# headers, so exact matches are claimed before any keyword is tried.
COLUMNS = {
    "location": ("footnote no./location in text", ("footnote", "location")),
    "checked": ("checked", ("checked",)),
    "reference": ("reference", ("reference",)),
    "quoted": ("quoted text", ("quoted",)),
    "discrepancy": ("discrepancy", ("discrepancy",)),
    "type": ("type of discrepancy", ("type",)),
    "severity": ("severity", ("severity",)),
    "description": ("description", ("description",)),
    "source": ("link/reference to the source you have consulted",
               ("link", "consulted")),
}

FILLS = {
    "none": "C6EFCE",          # discrepancy = no
    "low": "FFEB9C",
    "medium": "FCD5B4",
    "high": "FFC7CE",
    "unavailable": "E5E5E5",
}

METADATA_KEYS = ("document name", "date of source check",
                 "total references identified", "footnote references",
                 "main-text references")


def _cell(value):
    """A cell's text, with None and whitespace-only treated as empty."""
    if value is None:
        return ""
    return str(value).strip()


def _norm(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def _fill_of(cell):
    """The row-colour hex of a cell, or "" when it has no solid fill."""
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return ""
    rgb = getattr(fill.start_color, "rgb", None)
    if not isinstance(rgb, str):
        return ""
    return rgb[-6:].upper()


class Report:
    def __init__(self):
        self.violations = []   # things that are wrong
        self.notes = []        # things worth saying, not wrong
        self.stats = {}

    def bad(self, rule, row, message):
        self.violations.append({"rule": rule, "row": row, "message": message})

    def note(self, message):
        self.notes.append(message)


def find_columns(sheet, report):
    """Map the logical column names onto column indexes from the header row."""
    headers = {}
    for idx, cell in enumerate(sheet[1], start=1):
        text = _norm(cell.value)
        if text:
            headers[idx] = text
    found = {}
    taken = set()
    for name, (exact, _) in COLUMNS.items():
        for idx, text in headers.items():
            if idx not in taken and text == _norm(exact):
                found[name] = idx
                taken.add(idx)
                break
    for name, (_, keywords) in COLUMNS.items():
        if name in found:
            continue
        for idx, text in headers.items():
            if idx not in taken and all(k in text for k in keywords):
                found[name] = idx
                taken.add(idx)
                break
    missing = [n for n in COLUMNS if n not in found]
    if missing:
        report.bad("columns", 1,
                   "the header row has no column for: %s (found: %s)"
                   % (", ".join(missing), "; ".join(headers.values()) or "nothing"))
    return found


def read_rows(sheet, cols):
    """Every data row as a dict of logical name -> text, plus its fills."""
    rows = []
    for r in range(2, sheet.max_row + 1):
        values = {name: _cell(sheet.cell(row=r, column=idx).value)
                  for name, idx in cols.items()}
        if not any(values.values()):
            continue  # trailing blank row
        fills = {_fill_of(sheet.cell(row=r, column=idx)) for idx in cols.values()}
        rows.append({"row": r, "values": values, "fills": fills})
    return rows


def expected_fill(values):
    """The colour a row in this state must have, or None if unconstrained."""
    checked = values["checked"].lower()
    if checked == "source unavailable":
        return FILLS["unavailable"]
    if checked == "no":
        return None  # not checked yet: the skill does not prescribe a colour
    discrepancy = values["discrepancy"].lower()
    if discrepancy == "no":
        return FILLS["none"]
    if discrepancy == "yes":
        return FILLS.get(values["severity"].lower())
    return None


def check_rows(rows, report):
    seen = defaultdict(list)
    for entry in rows:
        r, v = entry["row"], entry["values"]
        checked = v["checked"].lower()
        discrepancy = v["discrepancy"].lower()
        severity = v["severity"].lower()

        # 1. the location column is the row's key
        if not v["location"]:
            report.bad("key", r, "no value in the location column")
        else:
            seen[v["location"]].append(r)

        # 2. checked is one of the three permitted values
        if checked not in CHECKED_VALUES:
            report.bad("checked", r, "checked = %r; must be one of %s"
                       % (v["checked"] or "", "/".join(CHECKED_VALUES)))
            continue

        # 3. a source nobody could read has no finding attached to it
        if checked == "source unavailable":
            for col in ("discrepancy", "type", "severity"):
                if v[col]:
                    report.bad("unavailable", r,
                               "checked = source unavailable, but %s = %r "
                               "(it must be empty)" % (col, v[col]))
            if not v["description"]:
                report.bad("unavailable", r,
                           "checked = source unavailable with no description; "
                           "say which tools were tried")
        elif checked == "yes":
            # 4. a checked row has a verdict
            if discrepancy not in ("yes", "no"):
                report.bad("discrepancy", r,
                           "checked = yes, but discrepancy = %r; must be yes or no"
                           % (v["discrepancy"] or ""))
            elif discrepancy == "yes":
                # 5. a finding is described, labelled and graded
                if not v["type"]:
                    report.bad("finding", r, "discrepancy = yes with no type of discrepancy")
                if severity not in SEVERITIES:
                    report.bad("finding", r,
                               "discrepancy = yes, but severity = %r; must be %s"
                               % (v["severity"] or "", "/".join(SEVERITIES)))
                if not v["description"]:
                    report.bad("finding", r, "discrepancy = yes with no description")
                if not v["source"]:
                    report.bad("finding", r,
                               "discrepancy = yes with no link/reference to the "
                               "source consulted")
            else:
                for col in ("type", "severity"):
                    if v[col]:
                        report.bad("no-finding", r,
                                   "discrepancy = no, but %s = %r (it must be empty)"
                                   % (col, v[col]))
                if not v["source"]:
                    report.bad("no-finding", r,
                               "checked = yes with no link/reference to the "
                               "source consulted")

        # 6. a quotation discrepancy quotes the text it is about
        if _norm(v["type"]) == _norm("Quotation not verbatim") and not v["quoted"]:
            report.bad("quotation", r,
                       'type of discrepancy = "Quotation not verbatim" with an '
                       "empty quoted text column")

        # 7. the colour says the same thing as the columns
        want = expected_fill(v)
        fills = {f for f in entry["fills"] if f}
        if want is not None:
            if len(fills) > 1:
                report.bad("colour", r, "the cells of the row have different "
                                        "fills (%s); expected %s throughout"
                           % (", ".join(sorted(fills)), want))
            elif not fills:
                report.bad("colour", r, "no row colour; expected %s" % want)
            elif fills != {want}:
                report.bad("colour", r, "row colour is %s; expected %s"
                           % (fills.pop(), want))

    for key, at in sorted(seen.items()):
        if len(at) > 1:
            report.bad("key", at[0],
                       "the location %r is used by %d rows (%s); each row needs "
                       "its own key - suffix them a, b, c ..."
                       % (key, len(at), ", ".join(str(a) for a in at)))


def check_metadata(book, rows, path, report, final):
    if "Metadata" not in book.sheetnames:
        report.bad("metadata", 0, "the workbook has no Metadata sheet")
        return
    sheet = book["Metadata"]
    meta = {}
    for r in range(1, sheet.max_row + 1):
        key = _norm(sheet.cell(row=r, column=1).value)
        if key:
            meta[key] = _cell(sheet.cell(row=r, column=2).value)
    for key in METADATA_KEYS:
        if _norm(key) not in meta:
            report.bad("metadata", 0, "the Metadata sheet has no %r row" % key)

    def as_int(key):
        raw = meta.get(_norm(key), "")
        m = re.search(r"-?\d+", raw)
        return int(m.group()) if m else None

    total = as_int("total references identified")
    notes_n = as_int("footnote references")
    text_n = as_int("main-text references")
    if total is not None and total != len(rows):
        report.bad("metadata", 0,
                   "Metadata says %d references, the sheet has %d rows"
                   % (total, len(rows)))
    if None not in (total, notes_n, text_n) and notes_n + text_n != total:
        report.bad("metadata", 0,
                   "%d footnote + %d main-text references is %d, not the %d "
                   "given as the total" % (notes_n, text_n, notes_n + text_n, total))

    name = meta.get(_norm("document name"), "")
    stem = Path(path).stem
    if name and stem.lower().startswith("kildesjekk_"):
        if _norm(Path(name).stem) != _norm(stem[len("kildesjekk_"):]):
            report.note("the file is %s but Metadata calls the document %r; "
                        "SKILL.md §2 asks for the two to match" % (Path(path).name, name))

    date = meta.get(_norm("date of source check"), "")
    if date and not re.match(r"^\d{4}-\d{2}-\d{2}", date):
        report.bad("metadata", 0,
                   "date of source check is %r; use ISO format (YYYY-MM-DD)" % date)
    if final and total is None:
        report.bad("metadata", 0, "no usable total in the Metadata sheet")


def summarise(rows, report, final):
    checked = Counter(r["values"]["checked"].lower() for r in rows)
    severity = Counter(r["values"]["severity"].lower() for r in rows
                       if r["values"]["discrepancy"].lower() == "yes")
    report.stats = {
        "rows": len(rows),
        "checked": {k: checked.get(k, 0) for k in CHECKED_VALUES},
        "discrepancies": sum(severity.values()),
        "by_severity": {k: severity.get(k, 0) for k in SEVERITIES},
    }
    if checked.get("no"):
        message = ("%d row(s) still have checked = no" % checked["no"])
        if final:
            report.bad("unchecked", 0, message + "; §16 requires this to be 0")
        else:
            report.note(message + " (still to do)")


def validate(path, *, final=False):
    report = Report()
    try:
        book = load_workbook(path)
    except Exception as e:  # noqa: BLE001 - any openpyxl failure is fatal here
        print("cannot read %s: %s" % (path, e), file=sys.stderr)
        raise SystemExit(2)
    sheet = book[book.sheetnames[0]]
    cols = find_columns(sheet, report)
    if len(cols) < len(COLUMNS):
        return report, []
    rows = read_rows(sheet, cols)
    check_rows(rows, report)
    check_metadata(book, rows, path, report, final)
    summarise(rows, report, final)
    return report, rows


def print_report(path, report):
    stats = report.stats
    print("%s - %d rows" % (Path(path).name, stats.get("rows", 0)))
    if stats:
        c = stats["checked"]
        print("  checked:       yes %d | source unavailable %d | no %d"
              % (c["yes"], c["source unavailable"], c["no"]))
        s = stats["by_severity"]
        print("  discrepancies: %d (high %d | medium %d | low %d)"
              % (stats["discrepancies"], s["high"], s["medium"], s["low"]))
    for n in report.notes:
        print("  note: %s" % n)
    if not report.violations:
        print("  no violations: the worklist is well formed.")
        print("  (This says nothing about whether the sources were really "
              "consulted or the quotations really compared.)")
        return
    print("  %d violation(s):" % len(report.violations))
    for v in report.violations:
        where = "row %d" % v["row"] if v["row"] else "workbook"
        print("    [%s] %s: %s" % (v["rule"], where, v["message"]))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    v = sub.add_parser("validate", help="check a worklist against the row invariants")
    v.add_argument("path")
    v.add_argument("--final", action="store_true",
                   help="also apply the §16 end-of-job rules (no row may be "
                        "left unchecked)")
    v.add_argument("--json", action="store_true", help="machine-readable output")
    args = parser.parse_args(argv)

    report, _ = validate(args.path, final=args.final)
    if args.json:
        print(json.dumps({"file": args.path, "stats": report.stats,
                          "violations": report.violations,
                          "notes": report.notes}, indent=2, ensure_ascii=False))
    else:
        print_report(args.path, report)
    return 1 if report.violations else 0


if __name__ == "__main__":
    sys.exit(main())
