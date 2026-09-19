#!/usr/bin/env python3
"""Tests for worklist.py validate.

Builds worklists with openpyxl in a temp directory - one well-formed, one
carrying the faults seen in real runs - and checks what the validator says
about them.

Run with `python plugins/kildesjekk/skills/kildesjekk/tests/test_worklist.py`.
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
except ImportError:
    sys.exit("these tests need openpyxl: pip install openpyxl")

import worklist  # noqa: E402

HEADERS = ["footnote no./location in text", "checked", "reference",
           "quoted text", "discrepancy", "type of discrepancy", "severity",
           "description", "link/reference to the source you have consulted"]

GREEN, YELLOW, ORANGE, RED, GREY = "C6EFCE", "FFEB9C", "FCD5B4", "FFC7CE", "E5E5E5"

failures = []


def check(name, got, want):
    if got == want:
        print("  ok   %s" % name)
    else:
        print("  FAIL %s: got %r, want %r" % (name, got, want))
        failures.append(name)


def rules_hit(report):
    return sorted({v["rule"] for v in report.violations})


def write_book(path, rows, metadata):
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    for values, colour in rows:
        sheet.append(values)
        if colour:
            for col in range(1, len(HEADERS) + 1):
                sheet.cell(row=sheet.max_row, column=col).fill = PatternFill(
                    fill_type="solid", start_color=colour)
    meta = book.create_sheet("Metadata")
    for key, value in metadata:
        meta.append([key, value])
    book.save(path)


CLEAN_ROWS = [
    (["1", "yes", "Eckhoff, Rettskildelære (2001) p. 23", "", "no", "", "",
      "", "Zotero ABCD1234"], GREEN),
    (["2a", "yes", "Case C-131/14 Google Spain, para 80",
      "«the data subject may request»", "yes", "Quotation not verbatim", "low",
      "The judgment reads «the data subject may, in the light of ...»",
      "CELEX:62014CJ0131"], YELLOW),
    (["2b", "source unavailable", "Müller, ZaöRV 2019 p. 44", "", "", "", "",
      "Not in Zotero; searched by author, title and a phrase from the quotation.",
      ""], GREY),
    (["text p. 22a", "yes", "Smith (1962) p. 149", "", "yes",
      "Wrong page/section/paragraph", "medium",
      "The passage is at p. 151, not p. 149.", "Zotero EFGH5678"], ORANGE),
    (["text p. 22b", "yes", "Jones (1978) p. 12", "", "yes", "Claim not supported",
      "high", "The source does not discuss the point at all.", "Zotero IJKL9012"], RED),
]

CLEAN_META = [
    ("Document name", "Artikkel.docx"),
    ("Date of source check", "2026-09-19"),
    ("Total references identified", 5),
    ("Footnote references", 3),
    ("Main-text references", 2),
]

with tempfile.TemporaryDirectory() as tmp:
    good = Path(tmp) / "Kildesjekk_Artikkel.xlsx"
    write_book(good, CLEAN_ROWS, CLEAN_META)

    print("a well-formed worklist")
    report, rows = worklist.validate(str(good))
    check("passes with no violations", report.violations, [])
    check("counts the rows", report.stats["rows"], 5)
    check("counts the statuses", report.stats["checked"],
          {"yes": 4, "no": 0, "source unavailable": 1})
    check("counts the discrepancies by severity", report.stats["by_severity"],
          {"high": 1, "medium": 1, "low": 1})

    print("a well-formed worklist with rows still to do")
    todo = Path(tmp) / "Kildesjekk_Todo.xlsx"
    write_book(todo, CLEAN_ROWS + [(["3", "no", "Hart (1961)", "", "", "", "", "", ""], None)],
               [("Document name", "Todo.docx"), ("Date of source check", "2026-09-19"),
                ("Total references identified", 6), ("Footnote references", 4),
                ("Main-text references", 2)])
    report, _ = worklist.validate(str(todo))
    check("an unchecked row is a note, not a violation", report.violations, [])
    check("and it is reported", any("checked = no" in n for n in report.notes), True)
    report, _ = worklist.validate(str(todo), final=True)
    check("--final turns it into a violation", rules_hit(report), ["unchecked"])

    # ----------------------------------------------------------------------
    # The faults from the September 2026 run: seven rows sharing one location
    # key, and "source unavailable" written into the discrepancy column.
    # ----------------------------------------------------------------------
    bad_rows = [
        (["1", "yes", "Eckhoff (2001) p. 23", "", "no", "", "", "", "Zotero ABCD"], GREEN),
        # discrepancy repeats the status instead of being left empty
        (["2", "source unavailable", "Müller, ZaöRV 2019", "", "source unavailable",
          "", "", "Not in Zotero.", ""], GREY),
        # ... and here it also picked up a severity and a colour to match
        (["3", "source unavailable", "Weber (2011)", "", "yes", "Claim not supported",
          "high", "Not in Zotero.", ""], RED),
        # duplicate keys: one results entry can only ever reach one of them
        (["text p.22 (§7, Recommended reading)", "yes", "First item", "", "no", "", "",
          "", "Zotero A1"], GREEN),
        (["text p.22 (§7, Recommended reading)", "yes", "Second item", "", "no", "", "",
          "", "Zotero A2"], GREEN),
        (["text p.22 (§7, Recommended reading)", "no", "Third item", "", "", "", "",
          "", ""], None),
        # a finding with no description, no severity and the wrong colour
        (["4", "yes", "Jones (1978) p. 12", "", "yes", "Claim not supported", "", "",
          "Zotero B1"], GREEN),
        # a quotation finding with nothing quoted
        (["5", "yes", "Smith (1962) p. 149", "", "yes", "Quotation not verbatim",
          "low", "Not verbatim.", "Zotero C1"], YELLOW),
        # a status value that is not one of the three
        (["6", "checked", "Hart (1961) p. 77", "", "no", "", "", "", "Zotero D1"], GREEN),
    ]
    bad = Path(tmp) / "Kildesjekk_Bad.xlsx"
    write_book(bad, bad_rows, [("Document name", "Bad.docx"),
                               ("Date of source check", "19.09.2026"),
                               ("Total references identified", 12),
                               ("Footnote references", 4),
                               ("Main-text references", 2)])

    print("a worklist with the faults seen in real runs")
    report, _ = worklist.validate(str(bad))

    def messages(rule):
        return [v["message"] for v in report.violations if v["rule"] == rule]

    check("the duplicated key is reported once, naming every row",
          [m for m in messages("key") if "3 rows" in m and "5, 6, 7" in m] != [], True)
    check("'source unavailable' in the discrepancy column is caught",
          any("discrepancy = 'source unavailable'" in m for m in messages("unavailable")), True)
    check("so is a severity on an unavailable row",
          len(messages("unavailable")) >= 3, True)
    check("a finding with no severity is caught",
          any("severity = ''" in m for m in messages("finding")), True)
    check("a finding with no description is caught",
          any("no description" in m for m in messages("finding")), True)
    check("an empty quoted text on a quotation finding is caught",
          len(messages("quotation")), 1)
    check("an unpermitted checked value is caught",
          any("checked = 'checked'" in m for m in messages("checked")), True)
    check("a row colour that contradicts the row is caught",
          any("expected" in m for m in messages("colour")), True)
    check("the Metadata row count mismatch is caught",
          any("the sheet has 9 rows" in m for m in messages("metadata")), True)
    check("the Metadata arithmetic is caught",
          any("is 6, not the 12" in m for m in messages("metadata")), True)
    check("a non-ISO date is caught",
          any("ISO format" in m for m in messages("metadata")), True)

    print("the CLI")
    check("exit 0 for a clean file", worklist.main(["validate", str(good)]), 0)
    check("exit 1 for a broken one", worklist.main(["validate", str(bad)]), 1)

    print("a workbook that is not a worklist")
    other = Path(tmp) / "Notes.xlsx"
    book = Workbook()
    book.active.append(["a", "b"])
    book.save(other)
    report, _ = worklist.validate(str(other))
    check("missing columns are reported, not crashed on", rules_hit(report), ["columns"])

print()
if failures:
    print("%d failing check(s): %s" % (len(failures), ", ".join(failures)))
    sys.exit(1)
print("all checks passed")
